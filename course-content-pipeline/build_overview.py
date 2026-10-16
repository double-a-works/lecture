import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

SRC = "Jadwal_TRA_Sistem.xlsx"
DAY_ORDER = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat"]

# (sheet_name, kelas_filter, title, batch_label, header_color, body_color)
BLOCKS = [
    ("Gasal 26-27 Sem1", "A", "SEM 1 KELAS A", "BATCH 2A", "1F4E78", "BDD7EE"),
    ("Gasal 26-27 Sem1", "B", "SEM 1 KELAS B", "BATCH 2B", "9C0006", "F4B6B6"),
    ("Gasal 26-27 Sem1", "C", "SEM 1 KELAS C", "BATCH 2C", "375623", "C6E0B4"),
    ("Gasal 26-27 Sem3", "Batch 1", "SEM 3", "BATCH 1", "833C00", "FCE4D6"),
]
# Only sheets within the same Tahun Akademik period are compared for conflicts.
PERIOD_SHEETS = ["Gasal 26-27 Sem1", "Gasal 26-27 Sem3"]
OUT_SHEET = "Overview Gasal 26-27"

wb = openpyxl.load_workbook(SRC, data_only=False)

# ---- 1. Read raw schedule rows from each Semester sheet in this period ----
all_entries = []  # dicts with semester,kelas,hari,jm,js,kode,mk,sks,dosen
for sn in PERIOD_SHEETS:
    ws = wb[sn]
    for r in range(2, ws.max_row + 1):
        hari = ws.cell(row=r, column=4).value
        if not hari:
            continue
        all_entries.append({
            "sheet": sn,
            "semester": ws.cell(row=r, column=2).value,
            "kelas": ws.cell(row=r, column=3).value,
            "hari": hari,
            "jm": ws.cell(row=r, column=5).value,
            "js": ws.cell(row=r, column=6).value,
            "kode": ws.cell(row=r, column=7).value,
            "mk": ws.cell(row=r, column=8).value,
            "sks": ws.cell(row=r, column=9).value,
            "dosen": ws.cell(row=r, column=10).value,
        })

# ---- 2. Conflict detection: same Hari + same Jam Mulai + same Dosen (non-blank) appears >1x ----
conflict_key_count = defaultdict(int)
for e in all_entries:
    if e["dosen"]:
        key = (e["hari"], e["jm"], e["dosen"])
        conflict_key_count[key] += 1


def is_conflict(e):
    if not e["dosen"]:
        return False
    return conflict_key_count[(e["hari"], e["jm"], e["dosen"])] > 1


# ---- 3. Remove existing Overview sheet if present, create fresh ----
if OUT_SHEET in wb.sheetnames:
    del wb[OUT_SHEET]
ws = wb.create_sheet(OUT_SHEET, index=1)
ws.sheet_properties.tabColor = "1F4E78"

thin = Side(style="thin", color="808080")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
yellow_fill = PatternFill("solid", fgColor="FFFF00")

SLOT_COLS = [4, 8, 12, 16]  # D, H, L, P (1-indexed column numbers)

row_cursor = 2
for sheet_name, kelas_filter, title, batch_label, header_hex, body_hex in BLOCKS:
    header_fill = PatternFill("solid", fgColor=header_hex)
    body_fill = PatternFill("solid", fgColor=body_hex)
    header_font = Font(bold=True, color="FFFFFF")

    entries = [e for e in all_entries if e["sheet"] == sheet_name and e["kelas"] == kelas_filter]
    by_day = defaultdict(list)
    for e in entries:
        by_day[e["hari"]].append(e)
    for d in by_day:
        by_day[d].sort(key=lambda e: str(e["jm"]))

    title_row = row_cursor
    sub_row = row_cursor + 1
    ws.cell(row=title_row, column=2, value=title)
    ws.merge_cells(start_row=title_row, start_column=2, end_row=title_row, end_column=3)
    ws.cell(row=title_row, column=4, value=batch_label)
    ws.merge_cells(start_row=title_row, start_column=4, end_row=title_row, end_column=5)
    for col in range(2, 20):
        c = ws.cell(row=title_row, column=col)
        c.fill = header_fill
        c.font = header_font
        c.border = border
    ws.cell(row=title_row, column=2).alignment = center
    ws.cell(row=title_row, column=4).alignment = center

    ws.cell(row=sub_row, column=2, value="No")
    ws.cell(row=sub_row, column=3, value="Hari")
    ws.cell(row=sub_row, column=4, value="Jam Perkuliahan")
    ws.merge_cells(start_row=sub_row, start_column=4, end_row=sub_row, end_column=19)
    for col in range(2, 20):
        c = ws.cell(row=sub_row, column=col)
        c.fill = header_fill
        c.font = header_font
        c.border = border
        c.alignment = center

    data_row = sub_row + 1
    for i, day in enumerate(DAY_ORDER, start=1):
        hr1 = data_row
        hr2 = data_row + 1
        ws.cell(row=hr1, column=2, value=i)
        ws.cell(row=hr1, column=3, value=day)
        ws.merge_cells(start_row=hr1, start_column=2, end_row=hr2, end_column=2)
        ws.merge_cells(start_row=hr1, start_column=3, end_row=hr2, end_column=3)
        for col in (2, 3):
            for rr in (hr1, hr2):
                c = ws.cell(row=rr, column=col)
                c.fill = body_fill
                c.border = border
                c.alignment = center

        day_entries = by_day.get(day, [])[:4]
        for slot_idx, base_col in enumerate(SLOT_COLS):
            for rr in (hr1, hr2):
                for cc in range(base_col, base_col + 4):
                    ws.cell(row=rr, column=cc).fill = body_fill
                    ws.cell(row=rr, column=cc).border = border
                    ws.cell(row=rr, column=cc).alignment = center
            if slot_idx >= len(day_entries):
                continue
            e = day_entries[slot_idx]
            ws.cell(row=hr1, column=base_col, value=e["jm"])
            ws.cell(row=hr1, column=base_col + 1, value=e["js"])
            ws.cell(row=hr1, column=base_col + 2, value="SKS")
            ws.cell(row=hr1, column=base_col + 3, value="Dosen")
            ws.merge_cells(start_row=hr2, start_column=base_col, end_row=hr2, end_column=base_col + 1)
            ws.cell(row=hr2, column=base_col, value=e["kode"])
            ws.cell(row=hr2, column=base_col + 2, value=e["sks"])
            ws.cell(row=hr2, column=base_col + 3, value=e["dosen"])
            if is_conflict(e):
                for cc in range(base_col, base_col + 4):
                    for rr in (hr1, hr2):
                        ws.cell(row=rr, column=cc).fill = yellow_fill
                ws.cell(row=hr2, column=base_col).font = Font(bold=True, color="CC0000")
                ws.cell(row=hr2, column=base_col + 3).font = Font(bold=True, color="CC0000")
        data_row += 2

    row_cursor = data_row + 1  # blank spacer row

# ---- 4. Column widths ----
ws.column_dimensions["A"].width = 3
ws.column_dimensions["B"].width = 5
ws.column_dimensions["C"].width = 9
for base in SLOT_COLS:
    ws.column_dimensions[get_column_letter(base)].width = 8
    ws.column_dimensions[get_column_letter(base + 1)].width = 8
    ws.column_dimensions[get_column_letter(base + 2)].width = 6
    ws.column_dimensions[get_column_letter(base + 3)].width = 9

ws.freeze_panes = "D2"
ws.sheet_view.showGridLines = False

note = (
    "Catatan: sheet ini dibuat otomatis dari data di sheet 'Gasal 26-27 Sem1' dan 'Gasal 26-27 Sem3'. "
    "Urutan slot per hari mengikuti urutan kronologis jam mulai (bukan posisi kolom tetap). "
    "Sel kuning = dosen bentrok jadwal (hari & jam mulai sama, terdeteksi lintas semester dalam periode Gasal 26-27)."
)
ws.cell(row=row_cursor, column=2, value=note)
ws.merge_cells(start_row=row_cursor, start_column=2, end_row=row_cursor, end_column=19)
ws.cell(row=row_cursor, column=2).font = Font(italic=True, size=9, color="666666")
ws.cell(row=row_cursor, column=2).alignment = Alignment(wrap_text=True, vertical="top")

wb.save(SRC)
print("Overview sheet built. Total entries:", len(all_entries), "Conflicts:", sum(1 for e in all_entries if is_conflict(e)))
